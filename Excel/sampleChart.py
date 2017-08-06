#coding=utf-8
#制作图表

import openpyxl
from openpyxl.chart import BarChart,Reference

wb = openpyxl.Workbook(write_only=True)
sheet = wb.create_sheet()

#图表数据
rows=[
	('Number','Batch1','Batch2'),
	(2,10,20),
	(3,20,30),
	(4,30,40),
	(5,40,50),
	(6,50,60),
	(7,60,70),
]
for i in rows:
	sheet.append(i)

#图表基本是设置
chart1=BarChart()
chart1.type='col'
chart1.style=11
chart1.title='Bar Chart'
chart1.y_axis.title='Test number'
chart1.x_axis.title='Sample length'


data = Reference(sheet,min_col=2,min_row=1,max_row=7,max_col=3)
cats = Reference(sheet,min_col=1,min_row=2,max_row=7)
chart1.add_data(data,titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 4
sheet.add_chart(chart1,'A11')

wb.save('sampleChart.xlsx')

